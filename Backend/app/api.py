from __future__ import annotations

import io
import json
import zipfile
from xml.sax.saxutils import escape
from pathlib import Path
from typing import Literal

from fastapi import FastAPI, Query, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

from .pipeline import PatentIdeaPipeline

DOMAIN_COLORS = {
    "agriculture": "#6ee7b7",
    "artificial-intelligence": "#7dd3fc",
    "chemical": "#f9a8d4",
    "sensors": "#c4b5fd",
}

DOMAIN_LABELS = {
    "agriculture": "Agriculture",
    "artificial-intelligence": "AI",
    "chemical": "Chemical",
    "sensors": "Sensors",
}


def _infer_domain(source_file: str | None) -> str:
    file_name = (source_file or "").lower()
    if "agriculture" in file_name:
        return "agriculture"
    if "artificial" in file_name or "ai" in file_name:
        return "artificial-intelligence"
    if "chemical" in file_name:
        return "chemical"
    if "sensor" in file_name:
        return "sensors"
    return "agriculture"


app = FastAPI(title="Invenza Research Observatory API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class ChatMessage(BaseModel):
    role: Literal["user", "assistant", "system"]
    content: str


class IdeaRequest(BaseModel):
    topic: str = Field(..., min_length=1)
    top_k: int = Field(default=3, ge=1, le=12)


class IdeaEvidence(BaseModel):
    patent_id: str
    title: str | None = None
    assignee: str | None = None
    abstract: str | None = None
    filing_date: str | None = None
    cosine_similarity: float = 0.0
    cpc_codes: list[str] = Field(default_factory=list)
    ipc_codes: list[str] = Field(default_factory=list)
    result_link: str | None = None
    representative_figure_link: str | None = None


class IdeaResponse(BaseModel):
    concept_name: str
    summary: str
    source_patents: list[str]
    used_patent_ids: list[str] = Field(default_factory=list)
    evidence: list[IdeaEvidence] = Field(default_factory=list)
    grounding_score: float
    passes_target: bool
    evidence_summary: str
    retrieval_details: dict[str, float]
    model_used: str = "llama3.2"
    inference_time_ms: float = 0.0
    total_latency_ms: float = 0.0


class ChatRequest(BaseModel):
    topic: str = Field(..., min_length=1)
    conversation: list[ChatMessage] = Field(default_factory=list)
    top_k: int = Field(default=3, ge=1, le=12)


class ChatResponse(BaseModel):
    response: str
    concept_name: str
    summary: str
    source_patents: list[str]
    used_patent_ids: list[str] = Field(default_factory=list)
    grounding_score: float
    evidence: list[IdeaEvidence] = Field(default_factory=list)
    evidence_summary: str = ""
    model_used: str = "llama3.2"
    inference_time_ms: float = 0.0
    conversation: list[ChatMessage]


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/patent-cluster")
def patent_cluster(limit: int = Query(default=300, ge=1, le=1500)) -> dict:
    dataset_path = Path(__file__).resolve().parents[1] / "data" / "patent_domain_data.jsonl"
    cluster_centers = {
        "agriculture": {"x": 220, "y": 115},
        "artificial-intelligence": {"x": 600, "y": 115},
        "chemical": {"x": 220, "y": 305},
        "sensors": {"x": 600, "y": 305},
    }
    nodes: list[dict[str, object]] = []
    cluster_counts: dict[str, int] = {key: 0 for key in DOMAIN_LABELS}
    per_domain_records: dict[str, list[dict[str, object]]] = {key: [] for key in DOMAIN_LABELS}

    if dataset_path.exists():
        with dataset_path.open("r", encoding="utf-8") as handle:
            for line in handle:
                if not line.strip():
                    continue
                try:
                    item = json.loads(line)
                except json.JSONDecodeError:
                    continue

                patent_id = item.get("patent_id") or item.get("metadata", {}).get("patent_id")
                if not patent_id:
                    continue

                source_file = item.get("source_file") or item.get("metadata", {}).get("source_file")
                domain = _infer_domain(source_file)
                per_domain_records.setdefault(domain, []).append(item)

    total_records = sum(len(items) for items in per_domain_records.values())
    sampled_domain_counts: dict[str, int] = {key: 0 for key in DOMAIN_LABELS}
    if total_records > 0:
        remaining = limit
        domains_order = list(DOMAIN_LABELS)
        for domain in domains_order:
            items = per_domain_records.get(domain, [])
            if not items:
                continue
            proportion = len(items) / total_records
            target = max(1, round(limit * proportion)) if remaining > 0 else 0
            sampled_domain_counts[domain] = min(len(items), target)
            remaining -= sampled_domain_counts[domain]
        if remaining > 0:
            for domain in domains_order:
                items = per_domain_records.get(domain, [])
                while remaining > 0 and sampled_domain_counts[domain] < len(items):
                    sampled_domain_counts[domain] += 1
                    remaining -= 1

    for domain in DOMAIN_LABELS:
        items = per_domain_records.get(domain, [])
        sample_count = sampled_domain_counts.get(domain, 0)
        for item in items[:sample_count]:
            center = cluster_centers.get(domain, {"x": 400, "y": 240})
            cluster_counts[domain] = cluster_counts.get(domain, 0) + 1
            index = cluster_counts[domain] - 1
            offset_x = (index % 13) * 12 - 72
            offset_y = (index // 13) * 10 - 45
            title = item.get("metadata", {}).get("title") or item.get("text", "").split("Title:", 1)[-1].split("\n", 1)[0] if item.get("text") else item.get("patent_id")
            nodes.append(
                {
                    "id": str(item.get("patent_id") or item.get("metadata", {}).get("patent_id")),
                    "title": str(title),
                    "domain": domain,
                    "cluster": DOMAIN_LABELS.get(domain, "Research"),
                    "color": DOMAIN_COLORS.get(domain, "#7dd3fc"),
                    "x": int(center["x"] + offset_x),
                    "y": int(center["y"] + offset_y),
                }
            )

    clusters = [
        {"id": key, "label": DOMAIN_LABELS.get(key, key), "color": DOMAIN_COLORS.get(key, "#7dd3fc"), "count": cluster_counts.get(key, 0)}
        for key in DOMAIN_LABELS
    ]

    return {"clusters": clusters, "nodes": nodes}


@app.post("/ideas", response_model=IdeaResponse)
def generate_idea(request: IdeaRequest) -> IdeaResponse:
    pipeline = PatentIdeaPipeline()
    result = pipeline.generate(request.topic, top_k=request.top_k)
    return IdeaResponse(**result)


@app.post("/chat", response_model=ChatResponse)
def chat(request: ChatRequest) -> ChatResponse:
    pipeline = PatentIdeaPipeline()
    conversation = [
        item.model_dump() if hasattr(item, "model_dump") else item
        for item in request.conversation
    ]
    latest_user_message = {"role": "user", "content": request.topic}
    if not conversation or conversation[-1] != latest_user_message:
        conversation.append(latest_user_message)
    result = pipeline.generate(
        request.topic,
        top_k=request.top_k,
        conversation=conversation,
    )

    assistant_message = {
        "role": "assistant",
        "content": (
            result["summary"].strip() or result["concept_name"]
        )
    }
    references = []
    reference_ids = set(result.get("used_patent_ids", []))
    for item in result.get("evidence", []):
        if item["patent_id"] in reference_ids and item["patent_id"] not in {
            reference["patent_id"] for reference in references
        }:
            references.append(item)
        if len(references) == 3:
            break
    if references:
        reference_text = "\n\nReference patents:\n" + "\n".join(
            f"- {item['patent_id']}: {item.get('title') or 'Untitled patent'}"
            + (f" ({item['result_link']})" if item.get("result_link") else "")
            for item in references
        )
        assistant_message["content"] += reference_text
    normalized_conversation = [*conversation, assistant_message]

    return ChatResponse(
        response=assistant_message["content"],
        concept_name=result["concept_name"],
        summary=result["summary"],
        source_patents=result["source_patents"],
        used_patent_ids=result.get("used_patent_ids", result["source_patents"]),
        grounding_score=result["grounding_score"],
        evidence=[IdeaEvidence(**item) for item in references],
        evidence_summary=result.get("evidence_summary", ""),
        model_used=result.get("model_used", "llama3.2"),
        inference_time_ms=result.get("inference_time_ms", 0.0),
        conversation=[ChatMessage(**item) for item in normalized_conversation],
    )


@app.post("/export/docx")
def export_docx(payload: dict) -> Response:
    concept_name = str(payload.get("concept_name") or "Patent idea")
    summary = str(payload.get("summary") or "")
    source = payload.get("source_patents") or []
    score = payload.get("grounding_score") or 0

    try:
        from docx import Document
    except ImportError:  # pragma: no cover - used only when python-docx is unavailable
        paragraphs = [concept_name, summary, f"Grounding score: {score}"]
        if source:
            paragraphs.extend(["Source patents", *[str(patent_id) for patent_id in source]])
        body = "".join(
            f'<w:p><w:r><w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p>'
            for text in paragraphs
        )
        document_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f"<w:body>{body}<w:sectPr/></w:body></w:document>"
        ).encode("utf-8")
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>""")
            archive.writestr("_rels/.rels", """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>""")
            archive.writestr("word/document.xml", document_xml)
    else:
        document = Document()
        document.add_heading(concept_name, level=0)
        document.add_paragraph(summary)
        document.add_paragraph(f"Grounding score: {score}")
        if source:
            document.add_heading("Source patents", level=1)
            for patent_id in source:
                document.add_paragraph(str(patent_id))

        buffer = io.BytesIO()
        document.save(buffer)

    buffer.seek(0)

    file_name = f"{concept_name.lower().replace(' ', '_')[:40]}.docx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{file_name}"'},
    )
